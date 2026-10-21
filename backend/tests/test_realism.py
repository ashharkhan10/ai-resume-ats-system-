"""
Tests for the "real ATS" checks: sections, layout, requirements, keyword placement,
exact skill matching, bullet quality, spelling checks and outcome bands.
"""

import io

from docx import Document

from app.schemas.resume import AIReading, LanguageIssue
from app.services.bullet_checks import check_bullets, quality_penalty
from app.services.keyword_match import match_keywords
from app.services.language import real_language_issues
from app.services.layout_check import inspect_layout
from app.services.requirements import check_requirements, score_requirements
from app.services.scoring.overall import band
from app.services.sections import detect_sections
from app.services.skills import mentioned_in, same_skill


# ---------- sections ----------

def test_standard_sections_are_found():
    text = "Priya\nEducation\nB.Tech\nSkills\nPython\nProjects\nChat app"
    result = detect_sections(text)
    assert result.found == ["Education", "Skills", "Projects"]
    assert result.missing == []


def test_missing_and_creative_sections_are_flagged():
    result = detect_sections("Priya\nMy Journey\nWorked at a cafe\nToolbox\nPython")
    assert "Education" in result.missing
    assert "Experience or Projects" in result.missing
    assert result.creative == ["My Journey", "Toolbox"]


# ---------- exact skill matching ----------

def test_short_skills_need_an_exact_match():
    assert not mentioned_in("C", "Wrote tools in C++ and C#")
    assert mentioned_in("C++", "Wrote tools in C++")
    assert not mentioned_in("Java", "Built a JavaScript app")
    assert not same_skill("Go", "Git")
    assert same_skill("ML", "Machine Learning")


# ---------- keyword placement ----------

def test_listed_only_skills_get_less_credit():
    result = match_keywords(
        resume_text="SKILLS Python, Docker\nBuilt an API in Python",
        resume_skills=["Python", "Docker"],
        required=["Python", "Docker"],
        preferred=[],
        used_text="Built an API in Python",
    )
    assert result.matched == ["Python", "Docker"]
    assert result.listed_only == ["Docker"]
    assert result.match_pct == 80  # (2 + 2×0.6) / 4


# ---------- requirements ----------

def _reading(**target) -> AIReading:
    return AIReading.model_validate({
        "resume": {"experience": [{"title": "Intern", "months": 6}], "certifications": ["AWS Cloud Practitioner"]},
        "target": target,
        "fit": {"education_met": False},
    })


def test_years_and_education_requirements():
    checks = check_requirements(_reading(min_years=2, education="B.Tech in CS"), "resume text")
    by_id = {c.id: c for c in checks}
    assert by_id["years"].ok is False
    assert "6 months" in by_id["years"].detail
    assert by_id["education"].ok is False
    assert score_requirements(checks) == 0


def test_certifications_are_found_on_the_resume():
    checks = check_requirements(_reading(certifications=["AWS Cloud Practitioner", "PMP"]), "text")
    assert [c.ok for c in checks] == [True, False]
    assert score_requirements(checks) == 50


def test_no_requirements_scores_full():
    assert score_requirements([]) == 100


# ---------- bullet quality ----------

def test_bullet_checks():
    findings = check_bullets([
        "I built a chat app with React",
        "Was responsible for testing the website",
        "Developed a website", "Developed an API", "Developed a bot",
    ])
    assert len(findings.first_person) == 1
    assert len(findings.passive) == 1
    assert findings.repeated_verbs == ["Developed ×3"]
    assert quality_penalty(findings) == 3 + 6 + 3


# ---------- spelling ----------

def test_made_up_language_issues_are_dropped():
    issues = [
        LanguageIssue(text="recieved", fix="received"),
        LanguageIssue(text="not in the resume", fix="whatever"),
        LanguageIssue(text="React", fix="React"),
    ]
    kept = real_language_issues(issues, "I recieved an award for my React app")
    assert [i.text for i in kept] == ["recieved"]


# ---------- layout (DOCX) ----------

def _docx_bytes(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_layout_problems_are_found():
    def build(document):
        document.sections[0].header.paragraphs[0].text = "priya@example.com | +91 98765 43210"
        document.add_paragraph("Experience")
        document.add_table(rows=2, cols=2)

    info = inspect_layout(_docx_bytes(build), "cv.docx")
    assert info.contact_in_header is True
    assert info.tables == 1
    assert info.two_columns is False


def test_broken_file_does_not_crash_layout_check():
    info = inspect_layout(b"not a real file", "cv.pdf")
    assert info.two_columns is None


# ---------- outcome bands ----------

def test_bands():
    assert band(80, True) == "Likely shortlisted"
    assert band(60, True) == "Borderline"
    assert band(30, True) == "Likely filtered out"
    assert band(80, False) == "Ready to apply"
