"""
Pulls plain text out of an uploaded resume (PDF or DOCX).
"""

import io

from docx import Document
from pypdf import PdfReader

from app.config import MAX_FILE_BYTES, MAX_FILE_MB

# A real resume has at least this many words of text.
MIN_WORDS = 40


class ResumeFileError(Exception):
    """A problem with the uploaded file. The message is shown to the user."""


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Checks the file and returns its text."""
    if not file_bytes:
        raise ResumeFileError("The file is empty.")
    if len(file_bytes) > MAX_FILE_BYTES:
        raise ResumeFileError(f"The file is bigger than {MAX_FILE_MB} MB.")

    name = filename.lower()
    if name.endswith(".pdf"):
        text = _read_pdf(file_bytes)
    elif name.endswith(".docx"):
        text = _read_docx(file_bytes)
    else:
        raise ResumeFileError("Please upload a PDF or DOCX file.")

    if len(text.split()) < MIN_WORDS:
        raise ResumeFileError(
            "We couldn't read enough text. If your resume is a scanned image, "
            "export it again as a text PDF from Word or Google Docs."
        )
    return text


def _read_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as error:
        raise ResumeFileError("This PDF couldn't be opened. Is it password-protected?") from error
    return "\n".join(pages).strip()


def _read_docx(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as error:
        raise ResumeFileError("This DOCX couldn't be opened. Try saving it again.") from error

    lines = [p.text for p in document.paragraphs if p.text.strip()]
    # Some resumes are built with tables; read those cells too.
    for table in document.tables:
        for row in table.rows:
            lines += [cell.text for cell in row.cells if cell.text.strip()]
    return "\n".join(lines).strip()
